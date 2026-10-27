"""
Lectura de proveedores desde Excel, validación de datos y escritura de resultados
(columnas Estado/Observaciones/Fecha por MKP + color de fila).
"""
import re
import shutil
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from config import BACKUP_DIR, COLORES, DATOS_INVALIDOS, OUTPUT_DIR

# Cada campo se reconoce porque el encabezado (sin acentos, minúsculas) empieza
# con alguno de estos alias. Así se toleran typos como "Correro electrónico".
ALIAS = {
    "nombre": ["nombre"],
    "correo": ["correo", "correro", "email", "e-mail"],
    "telefono": ["numero de telefono", "telefono", "celular"],
    "direccion": ["direccion", "plus code"],
    "cp": ["codigo postal", "cp", "c.p"],
    "pais": ["pais"],
    "descripcion": ["descripcion"],
    "disponible_24_7": ["disponibilidad", "24/7"],
    "tiempo_preparacion": ["tiempo de preparacion", "tiempo"],
}
OBLIGATORIOS = ["nombre", "correo", "telefono", "direccion", "cp"]

RE_CORREO = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
RE_PLUS_CODE = re.compile(r"\b[23456789CFGHJMPQRVWX]{4,8}\+[23456789CFGHJMPQRVWX]{2,3}\b", re.I)
VERDADEROS = {"true", "verdadero", "si", "sí", "1", "x", "yes"}


def norm(s) -> str:
    """minúsculas, sin acentos, espacios colapsados."""
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().lower()


def plus_code(direccion: str) -> str:
    m = RE_PLUS_CODE.search(direccion or "")
    return m.group(0).upper() if m else ""


@dataclass
class Proveedor:
    fila: int
    nombre: str = ""
    correo: str = ""
    telefono: str = ""
    direccion: str = ""
    cp: str = ""
    pais: str = "México"
    descripcion: str = ""
    disponible_24_7: bool = False
    tiempo_preparacion: int = 0
    errores: list = field(default_factory=list)

    @property
    def plus_code(self) -> str:
        return plus_code(self.direccion)


def _texto(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


class LibroProveedores:
    def __init__(self, ruta: Path, mkp: str):
        self.ruta = Path(ruta)
        self.mkp = mkp
        self.destino = self.ruta          # cambia si el Excel está abierto
        self.wb = load_workbook(self.ruta)
        self.ws = self.wb.worksheets[0]
        self.columnas = self._mapear_columnas()
        faltan = [c for c in OBLIGATORIOS if c not in self.columnas]
        if faltan:
            raise SystemExit(f"El Excel no tiene las columnas: {', '.join(faltan)}")
        self.ultima_col_datos = max(self.columnas.values())
        self.col_estado = self._col_resultado(f"Estado {mkp}")
        self.col_obs = self._col_resultado(f"Observaciones {mkp}")
        self.col_fecha = self._col_resultado(f"Fecha {mkp}")

    # ── columnas ──────────────────────────────
    def _mapear_columnas(self) -> dict:
        cols = {}
        for celda in self.ws[1]:
            h = norm(celda.value)
            if not h or h.startswith(("estado ", "observaciones ", "fecha ")):
                continue
            for campo, alias in ALIAS.items():
                if campo not in cols and any(h.startswith(a) for a in alias):
                    cols[campo] = celda.column
                    break
        return cols

    def _col_resultado(self, titulo: str) -> int:
        for celda in self.ws[1]:
            if norm(celda.value) == norm(titulo):
                return celda.column
        col = self.ws.max_column + 1
        c = self.ws.cell(row=1, column=col, value=titulo)
        c.font = Font(bold=True)
        self.ws.column_dimensions[c.column_letter].width = 45 if "Observ" in titulo else 18
        return col

    # ── lectura ───────────────────────────────
    def proveedores(self) -> list[Proveedor]:
        """Lee y valida todas las filas con datos. Las filas vacías se ignoran,
        así que se pueden agregar o quitar proveedores libremente."""
        lista, vistos = [], {}
        for r in range(2, self.ws.max_row + 1):
            crudo = {c: self.ws.cell(row=r, column=col).value for c, col in self.columnas.items()}
            if not any(_texto(v) for v in crudo.values()):
                continue
            p = Proveedor(fila=r)
            p.nombre = re.sub(r"\s+", " ", _texto(crudo.get("nombre")))
            p.correo = _texto(crudo.get("correo"))
            p.telefono = re.sub(r"\D", "", _texto(crudo.get("telefono")))
            p.direccion = _texto(crudo.get("direccion"))
            p.cp = _texto(crudo.get("cp"))
            if p.cp.isdigit() and len(p.cp) < 5:      # por si alguno vino como número
                p.cp = p.cp.zfill(5)
            p.pais = _texto(crudo.get("pais")) or "México"
            p.descripcion = _texto(crudo.get("descripcion"))
            p.disponible_24_7 = norm(_texto(crudo.get("disponible_24_7"))) in VERDADEROS \
                or crudo.get("disponible_24_7") is True
            t = _texto(crudo.get("tiempo_preparacion")) or "0"
            try:
                p.tiempo_preparacion = int(float(t))
                if p.tiempo_preparacion < 0:
                    raise ValueError
            except ValueError:
                p.errores.append(f"Tiempo de preparación inválido: '{t}'")

            if not p.nombre:
                p.errores.append("Sin nombre")
            if not RE_CORREO.match(p.correo):
                p.errores.append(f"Correo inválido: '{p.correo}'")
            if len(p.telefono) != 10:
                p.errores.append(f"Teléfono debe tener 10 dígitos: '{p.telefono}'")
            if not p.direccion:
                p.errores.append("Sin dirección")
            if not re.fullmatch(r"\d{5}", p.cp):
                p.errores.append(f"Código postal inválido: '{p.cp}'")
            clave = norm(p.nombre)
            if clave and clave in vistos:
                p.errores.append(f"Nombre duplicado en el Excel (fila {vistos[clave]})")
            vistos.setdefault(clave, r)
            lista.append(p)
        return lista

    def estado(self, fila: int) -> str:
        return _texto(self.ws.cell(row=fila, column=self.col_estado).value).upper()

    # ── escritura ─────────────────────────────
    def marcar(self, fila: int, estado: str, observaciones: str = "") -> None:
        self.ws.cell(row=fila, column=self.col_estado, value=estado)
        self.ws.cell(row=fila, column=self.col_obs, value=observaciones or None)
        self.ws.cell(row=fila, column=self.col_fecha, value=datetime.now().strftime("%Y-%m-%d %H:%M"))
        relleno = PatternFill("solid", fgColor=COLORES.get(estado, "FFFFFF"))
        for col in range(1, max(self.ultima_col_datos, self.col_fecha) + 1):
            self.ws.cell(row=fila, column=col).fill = relleno
        self.guardar()

    def respaldar(self) -> Path:
        BACKUP_DIR.mkdir(parents=True, exist_ok=True)
        dst = BACKUP_DIR / f"{self.ruta.stem}_{datetime.now():%Y%m%d_%H%M%S}{self.ruta.suffix}"
        shutil.copy2(self.ruta, dst)
        return dst

    def guardar(self) -> None:
        try:
            self.wb.save(self.destino)
        except PermissionError:
            # El Excel está abierto: se guarda una copia y se sigue usando esa.
            OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
            self.destino = OUTPUT_DIR / f"{self.ruta.stem}_resultado_{self.mkp}{self.ruta.suffix}"
            self.wb.save(self.destino)
