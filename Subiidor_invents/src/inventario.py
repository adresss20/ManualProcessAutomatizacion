"""
Archivos de inventario: ubicar el CSV que nombra el Excel, validarlo antes de subirlo
y (opcional) regenerarlo correctamente desde un .xlsx.

La validación existe porque Excel, al guardar un CSV, convierte los códigos de barras
en notación científica (7501349012004 -> 7.50135E+12). Un CSV así fusiona productos
distintos bajo el mismo SKU, por eso NO se sube salvo con --forzar.
"""
import csv
import re
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from config import INVENTARIOS_DIR

RE_CIENTIFICA = re.compile(r"^\s*-?\d+(?:[.,]\d+)?E[+-]?\d+\s*$", re.I)
COLUMNAS_CLAVE = ("SKU", "Title")


def _base(nombre: str) -> str:
    """'Inventario_tiendas.csv.' / 'Inventario_tiendas.csv..csv' -> 'inventario_tiendas'."""
    s = nombre.strip().lower()
    while True:
        nuevo = re.sub(r"(\.csv)?\.*$", "", s)
        if nuevo == s:
            return s
        s = nuevo


def resolver_archivo(nombre: str, carpeta: Path = INVENTARIOS_DIR):
    """Devuelve (ruta, nota). Primero el nombre exacto; si no existe, tolera puntos o
    '.csv' sobrantes/faltantes, siempre que haya UN solo candidato."""
    # Corrección mínima: quitar puntos/espacios sobrantes al final y asegurar ".csv".
    # (Ojo: Windows ignora el punto final, así que "x.csv." "existe" como "x.csv";
    # por eso se corrige el nombre ANTES de buscarlo.)
    corto = nombre.rstrip(". ")
    if not corto.lower().endswith(".csv"):
        corto += ".csv"
    if (carpeta / corto).is_file():
        return carpeta / corto, ("" if corto == nombre else f"En el Excel dice '{nombre}'; se usó '{corto}'")
    candidatos = [p for p in carpeta.glob("*.csv") if _base(p.name) == _base(nombre)]
    if len(candidatos) == 1:
        return candidatos[0], f"En el Excel dice '{nombre}'; se usó '{candidatos[0].name}'"
    if len(candidatos) > 1:
        return None, f"Hay varios CSV que podrían ser '{nombre}': {', '.join(p.name for p in candidatos)}"
    return None, f"No existe '{nombre}' en la carpeta Inventarios"


@dataclass
class ValidacionCSV:
    filas: int = 0
    skus_distintos: int = 0
    errores: list = field(default_factory=list)    # impiden subir
    avisos: list = field(default_factory=list)     # se sube, pero se anota

    @property
    def ok(self) -> bool:
        return not self.errores


def validar_csv(ruta: Path) -> ValidacionCSV:
    v = ValidacionCSV()
    if ruta.suffix.lower() != ".csv":
        v.errores.append(f"'{ruta.name}' no es .csv (el panel solo acepta CSV)")
        return v
    try:
        with open(ruta, encoding="utf-8-sig", newline="") as fh:
            filas = list(csv.reader(fh))
    except UnicodeDecodeError:
        v.errores.append("El CSV no está en UTF-8 (se verían mal los acentos). Guárdalo como 'CSV UTF-8'")
        return v
    if not filas:
        v.errores.append("El CSV está vacío")
        return v
    encabezado = [h.strip() for h in filas[0]]
    datos = [f for f in filas[1:] if any(c.strip() for c in f)]
    v.filas = len(datos)
    faltan = [c for c in COLUMNAS_CLAVE if c not in encabezado]
    if faltan:
        v.errores.append(f"Faltan columnas: {', '.join(faltan)}")
        return v
    i_sku = encabezado.index("SKU")
    skus = [f[i_sku].strip() if len(f) > i_sku else "" for f in datos]
    v.skus_distintos = len(set(s for s in skus if s))
    cientificos = sum(bool(RE_CIENTIFICA.match(s)) for s in skus)
    if cientificos:
        v.errores.append(f"{cientificos} de {v.filas} SKU están en notación científica (p. ej. '{next(s for s in skus if RE_CIENTIFICA.match(s))}'): "
                         "se perderían los códigos. Regenera el CSV con --convertir")
    vacios = sum(1 for s in skus if not s)
    if vacios:
        v.avisos.append(f"{vacios} filas sin SKU")
    if v.skus_distintos < v.filas - vacios:
        v.avisos.append(f"SKU repetidos: {v.filas - vacios - v.skus_distintos} filas")
    desalineadas = sum(1 for f in datos if len(f) != len(encabezado))
    if desalineadas:
        v.avisos.append(f"{desalineadas} filas con distinto número de columnas que el encabezado")
    return v


# ─────────────────────────────────────────────
#  Conversión xlsx -> csv (SKU como texto)
# ─────────────────────────────────────────────
def _celda(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "VERDADERO" if v else "FALSO"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return format(Decimal(repr(v)), "f")      # nunca notación científica
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return str(v)


def convertir_xlsx(xlsx: Path, destino: Path | None = None) -> Path:
    """Escribe un CSV UTF-8 (con BOM, como lo guarda Excel) con los valores reales de
    la primera hoja: los SKU quedan completos."""
    destino = destino or (INVENTARIOS_DIR / f"{Path(xlsx).stem}.csv")
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    with open(destino, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        for fila in ws.iter_rows(values_only=True):
            if fila and any(c is not None and str(c).strip() for c in fila):
                w.writerow([_celda(c) for c in fila])
    wb.close()
    return destino
